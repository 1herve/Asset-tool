"""Contrôles INVREF locaux ; aucune communication réseau ni correction des sources."""
from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import unicodedata
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Callable


# ============================================================
# 2. EXCEPTIONS PERSONNALISÉES
# ============================================================
# Ces erreurs permettent de distinguer :
# - une mauvaise configuration du référentiel INVREF ;
# - un problème dans le fichier de données fourni.
class ConfigurationError(ValueError):
    """Un référentiel incohérent ne peut pas être exécuté."""


class InputError(ValueError):
    """Erreur d'import explicite ; jamais convertie en fichier conforme."""


# ============================================================
# 3. FONCTIONS UTILITAIRES DE BASE
# ============================================================
# raw_text : transforme une valeur en texte de manière sûre.
# blank    : détermine si une cellule est réellement vide.
def raw_text(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def blank(value: Any) -> bool:
    # NaN, NULL, N/A sont des valeurs présentes : pas des vides silencieux.
    return value is None or (isinstance(value, str) and not value.strip())



# ============================================================
# 4. CONTRÔLE ET NORMALISATION DES NOMBRES
# ============================================================
# Cette fonction vérifie qu'une valeur est bien numérique.
# Elle contrôle notamment :
# - le séparateur décimal ;
# - le séparateur de milliers ;
# - les valeurs NaN / infini ;
# - les formats numériques incorrects.
# Elle renvoie un Decimal afin d'éviter les imprécisions des float.
def number(value: Any, settings: dict) -> Decimal:
    """Décimal fini, séparateurs explicitement configurés, groupement vérifié."""
    if isinstance(value, bool):
        raise ValueError('booléen au lieu de nombre')
    if isinstance(value, (int, float, Decimal)):
        result = Decimal(str(value))
    elif isinstance(value, str):
        text = value.strip()
        sep = settings['decimal_separator']
        group = settings['thousands_separator']
        if group == ' ':
            text = text.replace('\u00a0', ' ').replace('\u202f', ' ')
        decimal_part = re.escape(sep)
        integer = r'\d+'
        if group:
            integer = rf'(?:\d+|\d{{1,3}}(?:{re.escape(group)}\d{{3}})+)'
        pattern = rf'[+-]?(?:{integer}(?:{decimal_part}\d+)?|{decimal_part}\d+)(?:[eE][+-]?\d+)?'
        if not re.fullmatch(pattern, text, flags=re.ASCII):
            raise ValueError('syntaxe numérique ou séparateurs non conformes')
        if group:
            text = text.replace(group, '')
        result = Decimal(text.replace(sep, '.'))
    else:
        raise ValueError('type non numérique')
    if not result.is_finite():
        raise ValueError('NaN ou infini interdit')
    return result


# ============================================================
# 5. CONTRÔLE ET NORMALISATION DES DATES
# ============================================================
# Cette fonction vérifie qu'une date est valide.
# Elle accepte uniquement les formats définis dans la configuration.
# Elle sait également traiter les dates Excel lorsque cela est autorisé.
def parse_date(value: Any, settings: dict, epoch: Any = None) -> date:
    """Dates calendaires strictes ; numéro Excel seulement pour une cellule numérique."""
    if isinstance(value, datetime):
        if value.time() != datetime.min.time():
            raise ValueError('heure non nulle dans un champ date')
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, bool):
        raise ValueError('booléen au lieu de date')
    if isinstance(value, (int, float, Decimal)):
        if not math.isfinite(float(value)) or int(value) != value:
            raise ValueError('date numérique non entière')
        integer = int(value)
        if 10000000 <= integer <= 99999999:
            value = str(integer)
        elif settings['excel_serial_dates'] and epoch is not None:
            # Série 60 du calendrier Excel 1900 : 29/02/1900 n'existe pas.
            if epoch.year == 1899 and integer == 60:
                raise ValueError('jour fictif du calendrier Excel 1900')
            if integer <= 0:
                raise ValueError('numéro de date Excel non positif')
            from openpyxl.utils.datetime import from_excel
            parsed = from_excel(integer, epoch=epoch)
            if not isinstance(parsed, datetime):
                raise ValueError('numéro Excel ne représente pas un jour')
            return parsed.date()
        else:
            raise ValueError('numéro de date Excel interdit dans CSV ou hors paramétrage')
    if not isinstance(value, str):
        raise ValueError('type date incompatible')
    value = value.strip()
    for fmt in settings['date_formats']:
        try:
            parsed = datetime.strptime(value, fmt)
            # strptime accepte parfois un mois/jour sans zéro : imposer round trip.
            if parsed.strftime(fmt) == value:
                return parsed.date()
        except ValueError:
            continue
    raise ValueError('format ou jour calendaire invalide')

# ============================================================
# 6. CONTRÔLE DES IDENTIFIANTS ISIN
# ============================================================
# Vérifie :
# - la structure d'un code ISIN ;
# - sa longueur ;
# - son chiffre de contrôle grâce à l'algorithme de Luhn.
def isin_valid(value: str) -> bool:
    """Structure ISIN et chiffre de contrôle Luhn, seulement si explicitement activé."""
    if not re.fullmatch(r'[A-Z]{2}[A-Z0-9]{9}[0-9]', value, flags=re.ASCII):
        return False
    digits = ''.join(str(ord(c) - 55) if c.isalpha() else c for c in value)
    total = 0
    for i, c in enumerate(reversed(digits)):
        n = int(c) * (2 if i % 2 else 1)
        total += n // 10 + n % 10
    return total % 10 == 0

# ============================================================
# 7. CHARGEMENT DU RÉFÉRENTIEL DE CONTRÔLE
# ============================================================
# Charge le fichier invref.json contenant les règles des 98 variables.
# La configuration est ensuite validée avant utilisation.
def load_config(path: Path | None = None) -> dict:
    config_path = path or Path(__file__).with_name('invref.json')
    try:
        cfg = json.loads(config_path.read_text(encoding='utf-8'))
        validate_config(cfg)
    except (OSError, ValueError, KeyError, TypeError) as exc:
        raise ConfigurationError(f'Référentiel invalide : {exc}') from exc
    return cfg

# ============================================================
# 8. VALIDATION DE LA CONFIGURATION
# ============================================================
def validate_config(cfg: dict) -> None:
    fields = cfg['fields']
    if len(fields) != cfg['expected_field_count'] or len(fields) != 98:
        raise ConfigurationError('Le périmètre INVREF doit contenir exactement 98 variables.')
    for key in ('id', 'name'):
        if len({f[key] for f in fields}) != len(fields):
            raise ConfigurationError(f'{key} dupliqué')
    settings = cfg['settings']
    sep, group = settings['decimal_separator'], settings['thousands_separator']
    if sep not in ('.', ',') or group not in ('', ' ', '.', ',') or sep == group:
        raise ConfigurationError('Séparateurs numériques incompatibles')
    if not isinstance(settings['date_formats'], list) or not settings['date_formats']:
        raise ConfigurationError('Formats de dates manquants')
    for fmt in settings['date_formats']:
        if fmt not in ('%Y%m%d', '%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
            raise ConfigurationError('Format de date non pris en charge')
    for f in fields:
        if f['kind'] not in ('number', 'date', 'text', 'identifier', 'indicator'):
            raise ConfigurationError(f"Type inconnu : {f['name']}")
        if f['required'] is not None and type(f['required']) is not bool:
            raise ConfigurationError('required doit être true, false ou null')
        if f['case'] not in ('preserve', 'upper', 'lower'):
            raise ConfigurationError('Convention de casse invalide')
        if f['identifier_format'] not in (None, 'isin'):
            raise ConfigurationError('Format identifiant non pris en charge')
        if f['identifier_format'] and f['kind'] != 'identifier':
            raise ConfigurationError('ISIN doit être configuré sur un identifiant')
        if f['kind'] != 'number' and any(f[k] is not None for k in ('minimum','maximum')):
            raise ConfigurationError('Bornes numériques sur une variable non numérique')
        if f['kind'] != 'date' and any(f[k] is not None for k in ('date_min','date_max')):
            raise ConfigurationError('Bornes calendaires sur une variable non date')
        for p in f['pending']:
            if p['parameter'] not in f or not isinstance(p['reason'],str):
                raise ConfigurationError('Paramètre en attente inconnu ou motif invalide')
        for k in ('unique', 'exclusive_minimum', 'review_negative', 'review_nonpositive'):
            if type(f[k]) is not bool:
                raise ConfigurationError(f'{k} doit être booléen')
        if f['pattern'] is not None:
            try: re.compile(f['pattern'])
            except (re.error, TypeError) as exc: raise ConfigurationError('Expression régulière invalide') from exc
        if f['allowed_values'] is not None:
            if not isinstance(f['allowed_values'], list) or not f['allowed_values'] or any(not isinstance(x, str) for x in f['allowed_values']):
                raise ConfigurationError('allowed_values doit être une liste non vide de chaînes')
            if len(set(f['allowed_values'])) != len(f['allowed_values']):
                raise ConfigurationError('Valeurs autorisées dupliquées')
            canonical = [x.upper() if f['case']=='upper' else x.lower() if f['case']=='lower' else x for x in f['allowed_values']]
            if canonical != f['allowed_values'] or any(x.strip()!=x for x in canonical):
                raise ConfigurationError('Référentiel de valeurs non canonique')
        for k in ('min_length', 'max_length', 'max_distinct'):
            if f[k] is not None and (type(f[k]) is not int or f[k] < (1 if k=='max_distinct' else 0)):
                raise ConfigurationError(f'{k} doit être entier positif')
        if f['min_length'] is not None and f['max_length'] is not None and f['min_length'] > f['max_length']:
            raise ConfigurationError('Longueurs inversées')
        for k in ('minimum', 'maximum'):
            if f[k] is not None:
                try:
                    bound=Decimal(str(f[k]))
                except InvalidOperation as exc: raise ConfigurationError('Borne non numérique') from exc
                if not bound.is_finite(): raise ConfigurationError('Borne non finie')
        if f['minimum'] is not None and f['maximum'] is not None and Decimal(str(f['minimum'])) > Decimal(str(f['maximum'])):
            raise ConfigurationError('Bornes inversées')
        if f['exclusive_minimum'] and f['minimum'] is None:
            raise ConfigurationError('Minimum exclusif sans minimum')
        for k in ('date_min','date_max'):
            if f[k] is not None: date.fromisoformat(f[k])
        if f['date_min'] and f['date_max'] and f['date_min'] > f['date_max']:
            raise ConfigurationError('Bornes calendaires inversées')
        if f['max_blank_rate'] is not None and not 0 <= float(f['max_blank_rate']) <= 1:
            raise ConfigurationError('Seuil de vide hors [0,1]')

# ============================================================
# 9. LECTURE DES FICHIERS D'ENTRÉE CSV / EXCEL
# ============================================================
# Lit les données sans les modifier.
# Cette fonction :
# - accepte les fichiers .csv et .xlsx ;
# - récupère les en-têtes ;
# - conserve le numéro de ligne d'origine ;
# - détecte les formules Excel ;
# - gère les onglets et les lignes d'en-tête décalées.
def input_rows(path: Path, sheet: str | None, header_row: int, delimiter: str, encoding: str):
    """Conserve les en-têtes dupliqués et les lignes physiques (CSV multiligne inclus)."""
    if header_row < 1:
        raise InputError('header-row doit être >= 1')
    if path.suffix.lower() == '.csv':
        if len(delimiter) != 1:
            raise InputError('Un seul caractère séparateur attendu')
        with path.open(encoding=encoding, newline='') as stream:
            reader = csv.reader(stream, delimiter=delimiter, strict=True)
            for _ in range(header_row):
                try: headers = next(reader)
                except StopIteration as exc: raise InputError('En-tête introuvable') from exc
            yield ('header', headers, reader.line_num, None)
            while True:
                line = reader.line_num + 1
                try: values = next(reader)
                except StopIteration: break
                yield ('row', values, line, None)
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        # data_only=False : les formules sont détectées, jamais validées par cache.
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet is None and len(wb.sheetnames) != 1:
                raise InputError('Plusieurs onglets : préciser --sheet')
            if sheet is not None and sheet not in wb.sheetnames:
                raise InputError(f'Onglet absent : {sheet}')
            ws = wb[sheet or wb.sheetnames[0]]
            rows = ws.iter_rows(min_row=header_row)
            try: cells = next(rows)
            except StopIteration as exc: raise InputError('En-tête introuvable') from exc
            headers = [c.value for c in cells]
            while headers and headers[-1] is None: headers.pop()
            yield ('header', headers, header_row, wb.epoch)
            for line, cells in enumerate(rows, start=header_row+1):
                values = [c.value for c in cells]
                # Les cellules vides de mise en forme au-delà du schéma ne sont pas des colonnes.
                while len(values)>len(headers) and values[-1] is None: values.pop()
                if len(values)<len(headers): values += [None]*(len(headers)-len(values))
                formulas = [i for i,c in enumerate(cells) if c.data_type=='f']
                yield ('row', values, line, formulas)
        finally:
            wb.close()
    else:
        raise InputError('Formats pris en charge : .csv et .xlsx')

# ============================================================
# 10. PARAMÈTRES MÉTIER EN ATTENTE
# ============================================================
# Retourne les paramètres qui ne sont pas encore renseignés
# dans le référentiel pour une variable donnée.
# Exemple : borne min/max, domaine autorisé, format, etc.
def pending_parameters(f: dict) -> list[dict]:
    return [p for p in f['pending'] if f[p['parameter']] is None]

# ============================================================
# 11. LISTE DES CONTRÔLES ACTIFS POUR UNE VARIABLE
# ============================================================
# Détermine quels contrôles doivent être exécutés selon
# les paramètres renseignés dans invref.json.
def active_checks(f: dict, settings: dict) -> list[str]:
    checks = ['column','completeness','format','whitespace','control_characters']
    for key in ('allowed_values','pattern','min_length','max_length','minimum','maximum','date_min','date_max','max_distinct','max_blank_rate','identifier_format'):
        if f[key] is not None: checks.append(key)
    if f['required'] is True: checks.append('required')
    if f['unique']: checks.append('unique')
    if f['case'] != 'preserve': checks.append('case')
    if f['review_negative'] or f['review_nonpositive']: checks.append('review_sign')
    if (f['unique'] or f['max_distinct']) and not settings['track_distinct']:
        raise ConfigurationError('track_distinct doit être true pour les contrôles d’unicité')
    return checks

# ============================================================
# 12. MOTEUR PRINCIPAL D'AUDIT QDD
# ============================================================
# C'est la fonction centrale du programme.
#
# Elle réalise notamment :
# 1. la validation de la configuration ;
# 2. le contrôle du schéma (98 colonnes attendues) ;
# 3. les contrôles d'exhaustivité ;
# 4. les contrôles de conformité ;
# 5. la normalisation des valeurs ;
# 6. la collecte des anomalies ;
# 7. le calcul des statistiques par variable ;
# 8. la production d'une synthèse finale.
#
def audit(path: Path, cfg: dict, emit: Callable[[dict], None], *, sheet: str | None=None,
          header_row: int=1, delimiter: str=';', encoding: str='utf-8-sig',
          emit_normalized: Callable[[dict], None] | None=None) -> dict:
    """Émet les anomalies progressivement ; pas de conservation de toutes les lignes."""
    validate_config(cfg)
    fields = cfg['fields']; settings=cfg['settings']
    field_by_name = {f['name']:f for f in fields}
    stats = {f['name']:{'id':f['id'],'variable':f['name'],'present':False,'filled':0,'blank':0,'format_invalid':0,'anomaly_cells':0,'rules':active_checks(f,settings),'pending':pending_parameters(f),'deferred':f['deferred'],'counts':Counter()} for f in fields}
    distinct = {f['name']:set() for f in fields} if settings['track_distinct'] else {}
    counts=Counter(); rows=0; empty_rows=0; schema_errors=0
    checked_cells=0; blank_headers=[]
    def alert(rule: str, message: str, line: int | None=None, f: dict | None=None,
              raw: Any=None, normalized: Any=None, threshold: Any=None, severity: str | None=None):
        nonlocal schema_errors
        counts[rule]+=1
        if f is not None: stats[f['name']]['counts'][rule]+=1
        else: schema_errors+=1
        emit({'source_line':line,'field':f['name'] if f else '',
              'rule_id':f"{f['id'] if f else 'SCHEMA'}.{rule}",
              'severity':severity or (f['severity'] if f else 'Critique'),
              'raw_value':raw_text(raw),'normalized_value':raw_text(normalized),
              'threshold':raw_text(threshold),'message':message,'review_status':'à analyser'})
    iterator=iter(input_rows(path,sheet,header_row,delimiter,encoding))
    _, headers, header_line, epoch=next(iterator)
    normalized_headers=[]
    for i,h in enumerate(headers):
        text=raw_text(h)
        normalized=unicodedata.normalize('NFC',text).strip()
        if text!=normalized:
            alert('HEADER_NORMALIZATION','En-tête normalisé pour correspondance ; valeur brute conservée.',header_line,raw=text,normalized=normalized,severity='Élevée')
        if not normalized:
            blank_headers.append(i);alert('EMPTY_HEADER','Colonne sans nom.',header_line,raw=text)
        normalized_headers.append(normalized)
    duplicates={h for h,n in Counter(normalized_headers).items() if n>1}
    for h in sorted(duplicates):alert('DUPLICATE_HEADER','Nom de colonne dupliqué ; contrôles de cette colonne non exécutés.',header_line,raw=h)
    for h in normalized_headers:
        if h and h not in field_by_name:alert('UNEXPECTED_COLUMN','Colonne hors schéma INVREF.',header_line,raw=h)
    positions={h:i for i,h in enumerate(normalized_headers) if h in field_by_name and h not in duplicates}
    for f in fields:
        s=stats[f['name']];s['present']=f['name'] in positions
        if f['name'] not in normalized_headers:alert('MISSING_COLUMN','Colonne attendue absente.',header_line,f)
    if settings['check_order'] and normalized_headers != [f['name'] for f in fields]:
        alert('COLUMN_ORDER','Ordre de colonnes différent du schéma.',header_line)
    for _, values, line, metadata in iterator:
        # Une ligne CSV vide et une ligne Excel mise en forme ne sont pas une position.
        if not values or all(v is None or v=='' for v in values):
            empty_rows+=1;continue
        rows+=1
        if len(values)!=len(headers):
            alert('ROW_WIDTH','Nombre de cellules différent du nombre de colonnes.',line,raw=len(values),threshold=len(headers))
        formulas=set(metadata or [])
        normalized_row={'source_line':line}
        for f in fields:
            name=f['name'];s=stats[name]
            if name not in positions: continue
            idx=positions[name];v=values[idx] if idx<len(values) else None
            checked_cells+=1
            cell_count=sum(s['counts'].values())
            if blank(v):
                s['blank']+=1
                normalized_row[name]=''
                if f['required'] is True:alert('REQUIRED','Champ obligatoire vide.',line,f,v)
            else:
                s['filled']+=1
                text=raw_text(v);norm=text.strip()
                if isinstance(v,str) and text!=norm:
                    alert('WHITESPACE','Espaces en début ou fin de valeur.',line,f,v,norm)
                if any(unicodedata.category(c)=='Cc' for c in text):
                    alert('CONTROL_CHAR','Caractère de contrôle présent.',line,f,v,norm)
                if f['kind'] in ('identifier','indicator','text') and re.search(r' {2,}|[\u00a0\u202f]',norm):
                    alert('TEXT_SPACING','Espaces répétés ou insécables à examiner.',line,f,v,norm,severity='À analyser')
                try:
                    if idx in formulas:
                        raise ValueError('formule Excel non exécutée dans une extraction')
                    if f['kind']=='number': norm=number(v,settings)
                    elif f['kind']=='date': norm=parse_date(v,settings,epoch)
                    else:
                        if isinstance(v,bool) and f['kind']=='indicator':
                            norm='TRUE' if v else 'FALSE'
                        elif isinstance(v,(datetime,date,bool)):
                            raise ValueError('date ou booléen au lieu de code/texte')
                        if isinstance(v,(int,float,Decimal)) and not isinstance(v,bool):
                            if not math.isfinite(float(v)) or int(v)!=v:
                                raise ValueError('code numérique non entier / non fini')
                            norm=str(int(v))
                            alert('NUMERIC_CODE_STORAGE','Code stocké en nombre Excel : zéros initiaux potentiellement perdus.',line,f,v,norm,severity='À analyser')
                        if f['case']=='upper':norm=norm.upper()
                        if f['case']=='lower':norm=norm.lower()
                        if f['case']!='preserve' and text.strip()!=norm:
                            alert('CASE','Casse non canonique.',line,f,v,norm)
                        if not norm:raise ValueError('texte vide')
                    if f['allowed_values'] is not None and raw_text(norm) not in f['allowed_values']:
                        alert('DOMAIN','Valeur hors référentiel autorisé.',line,f,v,norm,f['allowed_values'])
                    if f['pattern'] is not None and not re.fullmatch(f['pattern'],raw_text(norm),flags=re.ASCII):
                        alert('PATTERN','Format non conforme au référentiel.',line,f,v,norm,f['pattern'])
                    if f['identifier_format']=='isin' and not isin_valid(raw_text(norm)):
                        alert('ISIN','Structure ou chiffre de contrôle ISIN invalide.',line,f,v,norm,'ISIN / Luhn')
                    for key, code, bad in [('min_length','MIN_LENGTH',lambda n,t:n<t),('max_length','MAX_LENGTH',lambda n,t:n>t)]:
                        if f[key] is not None and bad(len(raw_text(norm)),f[key]):alert(code,'Longueur non conforme.',line,f,v,norm,f[key])
                    if f['kind']=='number':
                        if f['minimum'] is not None:
                            low=Decimal(str(f['minimum']))
                            if norm<low or (f['exclusive_minimum'] and norm==low):alert('MINIMUM','Borne inférieure non respectée.',line,f,v,norm,('> ' if f['exclusive_minimum'] else '>= ')+str(low))
                        if f['maximum'] is not None and norm>Decimal(str(f['maximum'])):alert('MAXIMUM','Borne supérieure non respectée.',line,f,v,norm,f['maximum'])
                        if (f['review_nonpositive'] and norm<=0) or (f['review_negative'] and norm<0):alert('REVIEW_SIGN','Signe atypique à analyser ; pas de rejet métier automatique.',line,f,v,norm,severity='À analyser')
                    if f['kind']=='date':
                        if f['date_min'] and norm<date.fromisoformat(f['date_min']):alert('DATE_MIN','Date antérieure à la borne.',line,f,v,norm,f['date_min'])
                        if f['date_max'] and norm>date.fromisoformat(f['date_max']):alert('DATE_MAX','Date postérieure à la borne.',line,f,v,norm,f['date_max'])
                    if settings['track_distinct']:
                        key=norm if isinstance(norm,(Decimal,date)) else raw_text(norm);seen=distinct[name]
                        if f['unique'] and key in seen:alert('DUPLICATE_VALUE','Valeur dupliquée dans un champ configuré unique.',line,f,v,norm)
                        if f['max_distinct'] is not None and key not in seen and len(seen)>=f['max_distinct']:
                            alert('MAX_DISTINCT','Nombre maximal de valeurs distinctes dépassé.',line,f,v,norm,f['max_distinct'])
                        seen.add(key)
                    normalized_row[name]=raw_text(norm)
                except (ValueError, InvalidOperation, OverflowError) as exc:
                    s['format_invalid']+=1
                    normalized_row[name]=''
                    alert('FORMAT',str(exc),line,f,v)
            if sum(s['counts'].values())>cell_count: s['anomaly_cells']+=1
        if emit_normalized is not None:emit_normalized(normalized_row)
    if rows==0:alert('EMPTY_INPUT','Aucune ligne de données ; aucun résultat de conformité validable.')
    summary=[]
    for f in fields:
        s=stats[f['name']]
        s['blank_rate']=s['blank']/rows if rows and s['present'] else None
        s['distinct_values']=len(distinct[f['name']]) if settings['track_distinct'] and s['present'] else None
        if f['max_blank_rate'] is not None and s['blank_rate'] is not None and s['blank_rate']>f['max_blank_rate']:
            alert('BLANK_RATE','Taux de vide supérieur au seuil.',f=f,raw=s['blank_rate'],threshold=f['max_blank_rate'])
        s['status']='NON_EXECUTE' if not s['present'] or rows==0 else 'ANOMALIES' if s['counts'] else 'A_COMPLETER' if s['pending'] else 'CONFORME_CONTROLES_EXECUTES'
        s['counts']=dict(s['counts']);summary.append(s)
    return {'rows':rows,'empty_rows_skipped':empty_rows,'expected_variables':98,
            'present_variables':sum(s['present'] for s in summary),'checked_cells':checked_cells,
            'anomalies':sum(counts.values()),'schema_anomalies':schema_errors,
            'rules_counts':dict(counts),'pending_parameters':sum(len(s['pending']) for s in summary),
            'status':'NON_EXECUTE' if rows==0 else 'ANOMALIES' if counts else 'A_COMPLETER' if any(s['pending'] for s in summary) else 'CONFORME_CONTROLES_EXECUTES',
            'variables':summary}


def sha256_file(path: Path) -> str:
    digest=hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda:stream.read(1024*1024),b''):digest.update(block)
    return digest.hexdigest()
